"""エクスポート（CSV/Excel/PDF）のテスト（新スキーマ）。"""
import io

from openpyxl import load_workbook
from app.models import Department
from tests.conftest import login


def _dept2_id(app):
    with app.app_context():
        return Department.query.filter_by(name="第2営業部").first().id


def test_export_csv(client, app):
    login(client, "admin", "adminpass1")
    did = _dept2_id(app)
    resp = client.get(f"/projects/export.csv?dept={did}")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers["Content-Type"]
    text = resp.data.decode("utf-8-sig")
    assert "完成月" in text          # 新ヘッダー（計上月→完成月）
    assert "担当者" in text
    assert "取引先" in text          # 追加項目
    assert "商品カテゴリ" in text     # 追加項目
    assert "売上総利益" in text
    assert "計画案件A" in text       # conftest のデータ
    assert "1,000,000" in text       # 売上はカンマ表示


def test_export_xlsx(client, app):
    login(client, "admin", "adminpass1")
    did = _dept2_id(app)
    resp = client.get(f"/projects/export.xlsx?dept={did}")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["Content-Type"]
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "完成月"
    values = [c.value for row in ws.iter_rows(min_row=2) for c in row]
    assert "計画案件A" in values
    assert 1000000 in values         # 売上は数値として格納


def test_export_pdf(client, app):
    login(client, "admin", "adminpass1")
    did = _dept2_id(app)
    resp = client.get(f"/projects/export.pdf?dept={did}")
    assert resp.status_code == 200
    assert resp.headers["Content-Type"] == "application/pdf"
    assert resp.data[:4] == b"%PDF"
    assert len(resp.data) > 1000


def test_export_requires_login(client):
    resp = client.get("/projects/export.csv", follow_redirects=False)
    assert resp.status_code == 302
    assert "/login" in resp.headers["Location"]
