"""インポート（importers.py + imports ルート）と計画→案件管理コピーのテスト。"""
import io
import re

from openpyxl import load_workbook

from app import importers
from app.extensions import db
from app.models import Project, Department, Kubun, Rank, Category
from tests.conftest import login


# ---------- importers（HTTP非依存ロジック） ----------
def _maps(app):
    with app.app_context():
        dept = Department.query.filter_by(name="第2営業部").first()
        members = {u.name: u.user_id for u in dept.members}
        kubun = {k.name: k.id for k in Kubun.query.all()}
        cats = {}
        for c in Category.query.filter_by(department_id=dept.id).all():
            cats[c.code] = c.id
            cats[c.display_name] = c.id
        ranks = {r.name: r.id for r in Rank.query.all()}
        return members, kubun, cats, ranks


def test_validate_ok(app):
    members, kubun, cats, ranks = _maps(app)
    headers = importers.TEMPLATE_HEADERS
    rows = [(2, ["2026-09", "鈴木花子", "新規", "Ri", "取込案件", "○",
                 "1,000,000", "600000", "10", "8", "メモ"])]
    results, errors = importers.validate(headers, rows, members, kubun, cats, ranks)
    assert errors == []
    assert len(results) == 1
    d = results[0].data
    assert d["accounting_month"] == "2026-09"
    assert d["assignee_user_id"] == "hanako"
    assert d["project_name"] == "取込案件"
    assert d["sales"] == 1000000 and d["cost"] == 600000
    assert d["estimated_hours"] == 10.0


def test_validate_missing_required(app):
    members, kubun, cats, ranks = _maps(app)
    headers = importers.TEMPLATE_HEADERS
    # 計上月なし・案件名なし
    rows = [(2, ["", "鈴木花子", "新規", "Ri", "", "○", "100", "50", "", "", ""])]
    results, errors = importers.validate(headers, rows, members, kubun, cats, ranks)
    assert results == []
    assert any("計上月" in e for e in errors)
    assert any("案件名" in e for e in errors)


def test_validate_unknown_master(app):
    members, kubun, cats, ranks = _maps(app)
    headers = importers.TEMPLATE_HEADERS
    rows = [(2, ["2026-09", "存在しない人", "新規", "Ri", "案件", "○",
                 "100", "50", "", "", ""])]
    results, errors = importers.validate(headers, rows, members, kubun, cats, ranks)
    assert results == []
    assert any("担当者" in e for e in errors)


def test_template_xlsx_has_dropdown_and_format():
    choices = {"assignee": ["鈴木花子"], "kubun": ["期初計画", "新規"],
               "category": ["Ri", "Or"], "rank": ["○", "A", "B"]}
    data = importers.template_xlsx(choices)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert [c.value for c in ws[1]] == importers.TEMPLATE_HEADERS
    # ドロップダウン（データ検証）が設定されている
    assert len(ws.data_validations.dataValidation) >= 4


def test_template_xlsx_no_dropdown_when_empty():
    data = importers.template_xlsx(None)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert len(ws.data_validations.dataValidation) == 0


# ---------- ルート（アップロード→プレビュー→確定） ----------
def _dept_id(app):
    with app.app_context():
        return Department.query.filter_by(name="第2営業部").first().id


def test_template_download_dropdown_dept(client, app):
    login(client, "admin", "adminpass1")
    did = _dept_id(app)
    resp = client.get(f"/import/template.xlsx?dept={did}&type=initial")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["Content-Type"]
    wb = load_workbook(io.BytesIO(resp.data))
    # 第2営業部はドロップダウンあり
    assert len(wb.active.data_validations.dataValidation) >= 4


def test_import_roundtrip_initial(client, app):
    login(client, "admin", "adminpass1")
    did = _dept_id(app)
    # CSVを作成（計上月/担当者/区分/カテゴリー/案件名/確度/売上/仕入/見込み工数/対応工数/備考）
    csv_body = ("計上月,担当者,区分,カテゴリー,案件名,確度,売上,仕入,見込み工数,対応工数,備考\n"
                "2026-10,鈴木花子,新規,Or,インポート案件,A,\"2,000,000\",800000,,,取込\n")
    data = {"file": (io.BytesIO(csv_body.encode("utf-8-sig")), "up.csv")}
    resp = client.post(f"/import/preview?dept={did}&type=initial&period=59",
                       data=data, content_type="multipart/form-data",
                       follow_redirects=True)
    assert resp.status_code == 200
    assert "インポート案件".encode() in resp.data
    # プレビュー画面のトークンを取り出して確定
    m = re.search(rb'name="token" value="([0-9a-f]{32})"', resp.data)
    assert m, "確定用トークンが見つからない"
    token = m.group(1).decode()
    resp2 = client.post(f"/import/commit?dept={did}", data={"token": token},
                        follow_redirects=True)
    assert resp2.status_code == 200
    with app.app_context():
        p = Project.query.filter_by(project_name="インポート案件").first()
        assert p is not None
        assert p.plan_type == "initial"
        assert p.fiscal_period == 59
        assert p.sales == 2000000
        assert p.accounting_month == "2026-10"


def test_copy_plan_to_management(client, app):
    login(client, "admin", "adminpass1")
    did = _dept_id(app)
    with app.app_context():
        before_mgmt = Project.query.filter_by(
            department_id=did, fiscal_period=59, plan_type="management").count()
        initial_count = Project.query.filter_by(
            department_id=did, fiscal_period=59, plan_type="initial").count()
    assert initial_count >= 1
    resp = client.post(f"/projects/copy?dept={did}&period=59",
                       data={"source_type": "initial"}, follow_redirects=True)
    assert resp.status_code == 200
    with app.app_context():
        after_mgmt = Project.query.filter_by(
            department_id=did, fiscal_period=59, plan_type="management").count()
        # 既存の management は全置換され、initial の件数に一致する
        assert after_mgmt == initial_count
    assert before_mgmt  # 事前に management データがあったことを確認（全置換の意味づけ）
