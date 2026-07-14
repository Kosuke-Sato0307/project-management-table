"""案件データの CSV / Excel / PDF 出力。

列定義（COLUMNS）を1箇所に集約し、形式ごとに適切に整形する。
  - 金額: CSV/PDFはカンマ文字列、Excelは数値+桁区切り書式（合計可）
  - 日付: yyyy/m/d、完成月: yyyy/m（ExcelでもJul-26化しないよう文字列/書式で制御）
PDFの日本語は同梱フォント(ipaexg.ttf)を埋め込む（無ければ内蔵CIDフォントに自動フォールバック）。
"""
import csv
import io
from datetime import datetime, timezone, timedelta
from pathlib import Path

JST = timezone(timedelta(hours=9))

# (見出し, 属性, 種別)
COLUMNS = [
    ("計上月", "accounting_month", "month"),
    ("担当者", "assignee", "assignee"),
    ("区分", "kubun", "kubun"),
    ("カテゴリー", "category", "category"),
    ("案件名", "project_name", "text"),
    ("確度", "rank", "rank"),
    ("売上", "sales", "int"),
    ("仕入", "cost", "int"),
    ("売上総利益", "gross_profit", "int"),
    ("見込み工数", "estimated_hours", "float"),
    ("対応工数", "actual_hours", "float"),
    ("備考", "notes", "text"),
    ("作成日時", "created_at", "datetime"),
    ("作成者", "created_by", "text"),
    ("編集日時", "updated_at", "datetime"),
    ("編集者", "updated_by", "text"),
]

PDF_COLUMNS = ["計上月", "担当者", "区分", "カテゴリー", "案件名", "確度",
               "売上", "仕入", "売上総利益"]


# ---------- 値の取り出し・整形 ----------
def _raw(p, attr, kind):
    if kind == "assignee":
        return p.assignee.name if p.assignee else None
    if kind == "kubun":
        return p.kubun.name if p.kubun else None
    if kind == "category":
        return p.category.code if p.category else None
    if kind == "rank":
        return p.rank.name if p.rank else None
    return getattr(p, attr)


def _fmt_ymd(d):
    return f"{d.year}/{d.month}/{d.day}" if d else ""


def _fmt_month(s):
    if not s:
        return ""
    parts = str(s).replace("/", "-").split("-")
    if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit():
        return f"{int(parts[0])}/{int(parts[1])}"
    return str(s)


def _fmt_dt(v):
    if v is None:
        return ""
    if v.tzinfo is None:
        v = v.replace(tzinfo=timezone.utc)
    v = v.astimezone(JST)
    return f"{v.year}/{v.month}/{v.day} {v.strftime('%H:%M')}"


def _text_value(p, attr, kind):
    """CSV/PDF用: すべて文字列に整形。"""
    v = _raw(p, attr, kind)
    if kind == "int":
        return "" if v is None else f"{v:,}"
    if kind == "float":
        return "" if v is None else (f"{v:g}")
    if kind == "date":
        return _fmt_ymd(v)
    if kind == "month":
        return _fmt_month(v)
    if kind == "datetime":
        return _fmt_dt(v)
    return "" if v is None else str(v)


# ---------- CSV ----------
def to_csv(projects) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([h for h, _, _ in COLUMNS])
    for p in projects:
        writer.writerow([_text_value(p, attr, kind) for _, attr, kind in COLUMNS])
    return buf.getvalue().encode("utf-8-sig")


# ---------- Excel ----------
def to_xlsx(projects) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "案件一覧"

    ws.append([h for h, _, _ in COLUMNS])
    header_fill = PatternFill("solid", fgColor="2563EB")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for p in projects:
        row = []
        for _, attr, kind in COLUMNS:
            v = _raw(p, attr, kind)
            if kind in ("int", "float"):
                row.append(v)                       # 数値のまま
            elif kind == "date":
                row.append(v)                       # 実日付
            elif kind == "month":
                row.append(_fmt_month(v))           # 文字列 yyyy/m
            elif kind == "datetime":
                row.append(_fmt_dt(v))              # 文字列
            else:
                row.append("" if v is None else str(v))
        ws.append(row)

    ws.freeze_panes = "A2"
    for i, (h, _, kind) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(12, len(h) * 2 + 2)
        if kind in ("int", "float", "date", "month"):
            fmt = {"int": "#,##0", "float": "0.#", "date": "yyyy/m/d", "month": "@"}[kind]
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                for cell in row:
                    cell.number_format = fmt

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------- PDF ----------
_pdf_font_name = None


def _register_pdf_font() -> str:
    global _pdf_font_name
    if _pdf_font_name:
        return _pdf_font_name

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    font_path = Path(__file__).resolve().parent / "static" / "fonts" / "ipaexg.ttf"
    try:
        if font_path.exists():
            pdfmetrics.registerFont(TTFont("IPAexGothic", str(font_path)))
            _pdf_font_name = "IPAexGothic"
        else:
            raise FileNotFoundError
    except Exception:
        pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
        _pdf_font_name = "HeiseiKakuGo-W5"
    return _pdf_font_name


def to_pdf(projects, title="案件一覧") -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer)

    font = _register_pdf_font()
    styles = getSampleStyleSheet()
    jp = ParagraphStyle("jp", parent=styles["Normal"], fontName=font, fontSize=8, leading=11)
    jp_head = ParagraphStyle("jph", parent=jp, textColor=colors.white)
    title_style = ParagraphStyle("t", parent=styles["Title"], fontName=font, fontSize=14)
    small = ParagraphStyle("s", parent=jp, fontSize=8, textColor=colors.grey)

    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=landscape(A4),
                            leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm)

    col_def = {h: (attr, kind) for h, attr, kind in COLUMNS}
    header = [Paragraph(h, jp_head) for h in PDF_COLUMNS]
    data = [header]
    for p in projects:
        row = []
        for h in PDF_COLUMNS:
            attr, kind = col_def[h]
            row.append(Paragraph(_text_value(p, attr, kind), jp))
        data.append(row)

    total_w = 277 * mm
    weights = {"計上月": 0.9, "担当者": 1.1, "区分": 0.9, "カテゴリー": 0.8,
               "案件名": 2.4, "確度": 0.6, "売上": 1.2, "仕入": 1.2, "売上総利益": 1.2}
    wsum = sum(weights[h] for h in PDF_COLUMNS)
    col_widths = [total_w * weights[h] / wsum for h in PDF_COLUMNS]
    # 金額系の列は右寄せ
    num_headers = {"売上", "仕入", "売上総利益"}
    num_cols = [PDF_COLUMNS.index(h) for h in PDF_COLUMNS if h in num_headers]

    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for c in num_cols:
        style.append(("ALIGN", (c, 1), (c, -1), "RIGHT"))

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle(style))

    elems = [
        Paragraph(title, title_style),
        Paragraph(f"出力日時: {_fmt_dt(datetime.now(timezone.utc))}　件数: {len(projects)}", small),
        Spacer(1, 6),
        table,
    ]
    doc.build(elems)
    return bio.getvalue()
